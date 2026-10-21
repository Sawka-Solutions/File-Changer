import sys
import os
import tkinter as tk
from tkinter import filedialog
from PIL import Image
import imageio_ffmpeg
import subprocess
import fitz  # PyMuPDF
from pdf2docx import Converter
import docx
import webview

IMAGE_EXTS = ['.png', '.jpg', '.jpeg', '.webp', '.bmp', '.gif', '.tiff', '.ico']
AUDIO_EXTS = ['.mp3', '.wav', '.flac', '.aac', '.ogg', '.m4a']
VIDEO_EXTS = ['.mp4', '.avi', '.mkv', '.mov', '.webm', '.flv']
TEXT_EXTS = ['.txt', '.pdf', '.docx']

def get_category(ext):
    ext = ext.lower()
    if ext in IMAGE_EXTS: return 'image'
    if ext in AUDIO_EXTS: return 'audio'
    if ext in VIDEO_EXTS: return 'video'
    if ext in TEXT_EXTS: return 'text'
    return 'unknown'

def get_ffmpeg_path():
    return imageio_ffmpeg.get_ffmpeg_exe()

def choose_file_dialog():
    root = tk.Tk()
    root.withdraw()
    root.attributes('-topmost', True)
    file_path = filedialog.askopenfilename(
        title="WHAT ARE YOU STARTING WITH?",
        filetypes=[
            ("All supported files", "*.*"),
            ("Images", " ".join([f"*{x}" for x in IMAGE_EXTS])),
            ("Audio", " ".join([f"*{x}" for x in AUDIO_EXTS])),
            ("Video", " ".join([f"*{x}" for x in VIDEO_EXTS])),
            ("Text", " ".join([f"*{x}" for x in TEXT_EXTS]))
        ]
    )
    root.destroy()
    if not file_path:
        return None
    
    _, ext = os.path.splitext(file_path)
    category = get_category(ext)
    
    return {
        "path": file_path,
        "filename": os.path.basename(file_path),
        "category": category,
        "ext": ext.lower()
    }

def get_formats(category):
    if category == 'image': return [ext[1:] for ext in IMAGE_EXTS]
    if category == 'audio': return [ext[1:] for ext in AUDIO_EXTS]
    if category == 'video': return [ext[1:] for ext in VIDEO_EXTS]
    if category == 'text': return [ext[1:] for ext in TEXT_EXTS]
    return []

def do_convert_file(input_path, output_ext, preserve_formatting, save_special):
    try:
        _, ext = os.path.splitext(input_path)
        filename = os.path.basename(input_path)
        base_name = os.path.splitext(filename)[0]
        category = get_category(ext)
        
        # Ask for save path if save_special is true
        if save_special:
            root = tk.Tk()
            root.withdraw()
            root.attributes('-topmost', True)
            output_path = filedialog.asksaveasfilename(
                initialfile=f"{base_name}.{output_ext}",
                defaultextension=f".{output_ext}",
                title="WHERE TO?"
            )
            root.destroy()
            if not output_path:
                return {"success": False, "error": "Save cancelled"}
        else:
            output_path = os.path.join(os.path.dirname(input_path), f"{base_name}.{output_ext}")
            # Prevent overwriting the original file
            if os.path.abspath(output_path) == os.path.abspath(input_path):
                output_path = os.path.join(os.path.dirname(input_path), f"{base_name}_converted.{output_ext}")

        # --- CONVERSION LOGIC ---
        
        # 1. IMAGE
        if category == 'image':
            img = Image.open(input_path)
            
            if preserve_formatting == "ai":
                # Use AI to remove the background
                from rembg import remove
                img = remove(img)

            # convert rgba to rgb for jpeg/bmp
            if output_ext.lower() in ['jpg', 'jpeg', 'bmp'] and img.mode in ("RGBA", "P"):
                background = Image.new('RGB', img.size, (255, 255, 255))
                if img.mode == 'RGBA':
                    background.paste(img, mask=img.split()[3])
                else:
                    background.paste(img)
                img = background
                
            img.save(output_path)

        # 2. AUDIO & VIDEO
        elif category in ['audio', 'video']:
            ffmpeg_path = get_ffmpeg_path()
            cmd = [ffmpeg_path, "-y", "-i", input_path, output_path]
            subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)

        # 3. TEXT
        elif category == 'text':
            in_ext = ext.lower()
            out_ext = f".{output_ext.lower()}"
            
            if in_ext == out_ext:
                # Same format, just copy
                import shutil
                shutil.copy2(input_path, output_path)
            
            # PDF to TXT
            elif in_ext == '.pdf' and out_ext == '.txt':
                doc = fitz.open(input_path)
                text = ""
                for page in doc:
                    text += page.get_text()
                with open(output_path, "w", encoding="utf-8") as f:
                    f.write(text)
            
            # PDF to DOCX
            elif in_ext == '.pdf' and out_ext == '.docx':
                if preserve_formatting == "yes":
                    cv = Converter(input_path)
                    cv.convert(output_path)
                    cv.close()
                else:
                    # Just extract text to docx
                    doc = fitz.open(input_path)
                    new_doc = docx.Document()
                    for page in doc:
                        new_doc.add_paragraph(page.get_text())
                    new_doc.save(output_path)
            
            # DOCX to TXT
            elif in_ext == '.docx' and out_ext == '.txt':
                doc_obj = docx.Document(input_path)
                with open(output_path, "w", encoding="utf-8") as f:
                    for para in doc_obj.paragraphs:
                        f.write(para.text + "\n")
            
            # TXT to DOCX
            elif in_ext == '.txt' and out_ext == '.docx':
                with open(input_path, "r", encoding="utf-8") as f:
                    lines = f.readlines()
                new_doc = docx.Document()
                for line in lines:
                    new_doc.add_paragraph(line.strip())
                new_doc.save(output_path)
            
            # TXT to PDF
            elif in_ext == '.txt' and out_ext == '.pdf':
                # very simple text to pdf using fitz
                doc = fitz.open()
                page = doc.new_page()
                with open(input_path, "r", encoding="utf-8") as f:
                    text = f.read()
                # Insert text
                page.insert_text((50, 50), text, fontsize=11)
                doc.save(output_path)
            
            # DOCX to PDF
            elif in_ext == '.docx' and out_ext == '.pdf':
                if preserve_formatting == "yes":
                    try:
                        import win32com.client
                        word = win32com.client.Dispatch('Word.Application')
                        doc = word.Documents.Open(os.path.abspath(input_path))
                        doc.SaveAs(os.path.abspath(output_path), FileFormat=17)
                        doc.Close()
                        word.Quit()
                    except Exception as e:
                        return {"success": False, "error": f"Formatting preservation for DOCX to PDF requires Microsoft Word installed. Try 'No' formatting. Error: {str(e)}"}
                else:
                    doc_obj = docx.Document(input_path)
                    text = "\n".join([para.text for para in doc_obj.paragraphs])
                    out_doc = fitz.open()
                    page = out_doc.new_page()
                    page.insert_text((50, 50), text[:5000], fontsize=11)
                    out_doc.save(output_path)
            
            else:
                return {"success": False, "error": f"Conversion from {in_ext} to {out_ext} is not supported yet."}

        return {"success": True, "path": output_path}
    except Exception as e:
        return {"success": False, "error": str(e)}

class AppApi:
    def choose_file(self):
        return choose_file_dialog()

    def get_output_formats(self, category):
        return get_formats(category)

    def convert_file(self, input_path, output_ext, preserve_formatting, save_special):
        return do_convert_file(input_path, output_ext, preserve_formatting, save_special)

def get_resource_path(relative_path):
    if getattr(sys, 'frozen', False):
        base_dir = getattr(sys, '_MEIPASS', os.path.dirname(sys.executable))
    else:
        base_dir = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_dir, relative_path)

if __name__ == '__main__':
    api = AppApi()
    html_file = get_resource_path(os.path.join('web', 'index.html'))
    icon_file = get_resource_path('icon.ico')
    
    window = webview.create_window(
        title='File Changer',
        url=html_file,
        js_api=api,
        width=720,
        height=640,
        resizable=False,
        background_color='#0f1115'
    )
    webview.start(icon=icon_file if os.path.exists(icon_file) else None)
